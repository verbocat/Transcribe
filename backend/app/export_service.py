import json
import csv
import io
from pathlib import Path
from typing import List, Dict, Any, Optional
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.models import TranscriptionResult, Segment, AudioAnalysis
from app.config import EXPORTS_DIR


_TABLE_HEADERS = [
    "Segment ID",
    "Speaker",
    "Gender",
    "Start Time",
    "End Time",
    "Duration",
    "Transcript (Full Verbatim)",
]


def _has_translation(result: TranscriptionResult) -> bool:
    return any(seg.translation for seg in result.segments)


def _headers(result: TranscriptionResult) -> list:
    """Table headers; a second-language column is appended when the export carries a translation."""
    if not _has_translation(result):
        return list(_TABLE_HEADERS)
    label = f" ({result.translation_language})" if result.translation_language else ""
    return _TABLE_HEADERS + [f"Translation{label}"]


def _segment_row(seg: Segment, with_translation: bool = False) -> list:
    row = [
        seg.segment_id,
        seg.speaker,
        seg.gender,
        _hms_ms(seg.start_time),
        _hms_ms(seg.end_time),
        _hms_ms(seg.end_time - seg.start_time),
        seg.transcript,
    ]
    if with_translation:
        row.append(seg.translation or "")
    return row


def _hms_ms(seconds: float) -> str:
    """Seconds -> HH:MM:SS.mmm (e.g. 00:00:12.420)."""
    total_ms = max(0, int(round(float(seconds or 0.0) * 1000)))
    secs, ms = divmod(total_ms, 1000)
    return f"{secs // 3600:02d}:{(secs % 3600) // 60:02d}:{secs % 60:02d}.{ms:03d}"


def export_to_csv(result: TranscriptionResult, delimiter: str = ",") -> str:
    """Generate CSV or TSV string representation."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=delimiter)
    
    tr = _has_translation(result)
    writer.writerow(_headers(result))
    for seg in result.segments:
        writer.writerow(_segment_row(seg, tr))

    return output.getvalue()


def export_to_tsv(result: TranscriptionResult, delimiter: str = "\t") -> str:
    """Generate TSV string representation."""
    return export_to_csv(result, delimiter=delimiter)


def export_to_txt(result: TranscriptionResult) -> str:
    """Generate clean human-readable TXT transcript."""
    dur = result.audio_info.duration if result.audio_info else (result.segments[-1].end_time if result.segments else 0.0)
    lines = []
    lines.append(f"================================================================================")
    lines.append(f"LOWER THIRD TRANSCRIPTION DELIVERABLE: {result.filename}")
    lines.append(f"Language: {result.language} | Script: {result.script}")
    lines.append(f"Duration: {_hms_ms(dur)} | Compliance Score: {result.compliance_score}%")
    lines.append(f"================================================================================\n")

    for seg in result.segments:
        lines.append(f"[{seg.start_time_str} --> {seg.end_time_str}] {seg.speaker} ({seg.gender}):")
        lines.append(f"   {seg.transcript}")
        if seg.translation:
            lines.append(f"   {seg.translation}")
        lines.append("")

    return "\n".join(lines)


def export_to_srt(result: TranscriptionResult) -> str:
    """Generate standard SRT subtitle file."""
    def srt_time(seconds: float) -> str:
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        return f"{hrs:02d}:{mins:02d}:{secs:02d},{millis:03d}"

    entries = []
    for i, seg in enumerate(result.segments, 1):
        s_str = srt_time(seg.start_time)
        e_str = srt_time(seg.end_time)
        # SRT standard: subtitle body should be just the transcript text.
        # Speaker label is placed as a clean prefix (compatible with most SRT players).
        speaker_prefix = f"[{seg.speaker}] " if seg.speaker else ""
        body = f"{speaker_prefix}{seg.transcript}" + (f"\n{seg.translation}" if seg.translation else "")
        entries.append(f"{i}\n{s_str} --> {e_str}\n{body}\n")

    return "\n".join(entries)


def export_to_vtt(result: TranscriptionResult) -> str:
    """Generate standard WebVTT subtitle file (.vtt)."""
    def vtt_time(seconds: float) -> str:
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        return f"{hrs:02d}:{mins:02d}:{secs:02d}.{millis:03d}"

    lines = ["WEBVTT", ""]
    for i, seg in enumerate(result.segments, 1):
        s_str = vtt_time(seg.start_time)
        e_str = vtt_time(seg.end_time)
        speaker_prefix = f"[{seg.speaker}] " if seg.speaker else ""
        lines.append(f"{i}")
        lines.append(f"{s_str} --> {e_str}")
        lines.append(f"{speaker_prefix}{seg.transcript}")
        if seg.translation:
            lines.append(seg.translation)
        lines.append("")

    return "\n".join(lines)


def export_to_json(result: TranscriptionResult) -> str:
    """Generate clean segment list JSON matching exact Karya user deliverable format."""
    items = []
    for seg in result.segments:
        item = {
            "start_sec": round(float(seg.start_time), 3),
            "end_sec": round(float(seg.end_time), 3),
            "transcription": seg.transcript or "",
            "speaker": seg.speaker,
            "gender_label": seg.gender
        }
        if seg.translation:
            item["translation"] = seg.translation
        items.append(item)
    return json.dumps(items, ensure_ascii=False, indent=2)


def export_to_docx(result: TranscriptionResult, output_path: str) -> str:
    """Generate Word (.docx) transcript deliverable."""
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    FONT = "Calibri"
    INDIC_FONT = "Nirmala UI"  # complex-script font so Devanagari etc. render properly in Word

    def style_run(run, size, bold=False):
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        rpr = run._element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi"):
            fonts.set(qn(attr), FONT)
        fonts.set(qn("w:cs"), INDIC_FONT)

    def set_cell(cell, text, size, bold=False, align=None):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if align is not None:
            p.alignment = align
        style_run(p.add_run(str(text)), size, bold)

    def shade(cell, hex_fill):
        tc_pr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.7))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(title.add_run("Lower Third Transcription Deliverable"), 20, bold=True)

    dur = result.audio_info.duration if result.audio_info else (result.segments[-1].end_time if result.segments else 0.0)
    meta_lines = [
        ("Audio File: ", result.filename),
        ("Language & Script: ", f"{result.language} ({result.script})"),
        ("Duration: ", _hms_ms(dur)),
        ("Total Segments: ", str(len(result.segments))),
    ]
    for label, value in meta_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        style_run(p.add_run(label), 10.5, bold=True)
        style_run(p.add_run(value), 10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    tr = _has_translation(result)
    headers = _headers(result)
    widths = [Inches(0.8), Inches(1.0), Inches(0.8), Inches(1.15), Inches(1.15), Inches(1.15), Inches(3.55)]
    if tr:  # split the text column between the transcript and its translation
        widths = widths[:6] + [Inches(1.9), Inches(1.9)]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for i, title_text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell(cell, title_text.replace(" (Full Verbatim)", ""), 10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, "D9E2F3")
    # repeat the header row on every page
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    for seg in result.segments:
        cells = table.add_row().cells
        for i, value in enumerate(_segment_row(seg, tr)):
            last = i >= 6
            set_cell(cells[i], value, 10, align=None if last else WD_ALIGN_PARAGRAPH.CENTER)

    for col, width in zip(table.columns, widths):
        col.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    doc.save(output_path)
    return output_path


def export_to_xlsx(result: TranscriptionResult, output_path: str) -> str:
    """Generate formatted Excel (.xlsx) deliverable with summary and data sheets."""
    from openpyxl.styles import Alignment, Font, PatternFill

    dur = result.audio_info.duration if result.audio_info else (result.segments[-1].end_time if result.segments else 0.0)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        # Sheet 1: Data
        tr = _has_translation(result)
        df_data = pd.DataFrame([_segment_row(seg, tr) for seg in result.segments], columns=_headers(result))
        df_data.to_excel(writer, sheet_name="Transcription", index=False)

        # Sheet 2: Summary
        df_summary = pd.DataFrame([
            {"Metric": "Audio Filename", "Value": result.filename},
            {"Metric": "Language", "Value": result.language},
            {"Metric": "Script", "Value": result.script},
            {"Metric": "Duration", "Value": _hms_ms(dur)},
            {"Metric": "Total Segments", "Value": len(result.segments)},
            {"Metric": "Compliance Score (%)", "Value": result.compliance_score},
            {"Metric": "Total QC Errors", "Value": result.total_errors},
            {"Metric": "Total QC Warnings", "Value": result.total_warnings},
            {"Metric": "Is Rejected", "Value": "Yes" if result.is_rejected else "No"},
            {"Metric": "Rejection Reason", "Value": result.rejection_reason or "N/A"},
        ])
        df_summary.to_excel(writer, sheet_name="Audit Summary", index=False)

        # Formatting: bold shaded header, column widths, wrapped transcript, frozen header row
        header_fill = PatternFill("solid", fgColor="D9E2F3")
        ws = writer.sheets["Transcription"]
        for col, width in zip("ABCDEFGH", [11, 14, 10, 14, 14, 14, 90, 90] if tr else [11, 14, 10, 14, 14, 14, 90]):
            ws.column_dimensions[col].width = width
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=cell.column >= 7,
                                           horizontal=None if cell.column >= 7 else "center")
        ws.freeze_panes = "A2"
        ws2 = writer.sheets["Audit Summary"]
        ws2.column_dimensions["A"].width = 24
        ws2.column_dimensions["B"].width = 40
        for cell in ws2[1]:
            cell.font, cell.fill = Font(bold=True), header_fill

    return output_path


def export_rejection_csv(rejected_items: List[Dict[str, Any]]) -> str:
    """Generate Rejection Report CSV according to Karya Guideline 3."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Audio Filename", "Rejection Category", "Detailed Rejection Reason", "Duration (s)", "RMS (dBFS)", "SNR (dB)"])

    for item in rejected_items:
        writer.writerow([
            item.get("filename", ""),
            item.get("rejection_category", ""),
            item.get("rejection_reason", ""),
            item.get("duration", 0.0),
            item.get("rms_db", 0.0),
            item.get("snr_db", 0.0)
        ])

    return output.getvalue()


def _prepare_sorted_events(events: list) -> list:
    """Sort events chronologically and clamp any overlapping end times before export."""
    if not events:
        return []
    def _get_st(ev):
        val = ev.get("start_time", ev.get("start", 0.0)) if isinstance(ev, dict) else getattr(ev, "start_time", getattr(ev, "start", 0.0))
        return float(val or 0.0)
    sorted_evs = sorted(events, key=_get_st)
    for i in range(len(sorted_evs) - 1):
        cur = sorted_evs[i]
        nxt = sorted_evs[i + 1]
        c_et = float(cur.get("end_time", cur.get("end", 0.0)) if isinstance(cur, dict) else getattr(cur, "end_time", getattr(cur, "end", 0.0)))
        n_st = float(nxt.get("start_time", nxt.get("start", 0.0)) if isinstance(nxt, dict) else getattr(nxt, "start_time", getattr(nxt, "start", 0.0)))
        if c_et > n_st - 0.04:
            new_et = round(max(n_st - 0.04, 0.0), 3)
            if isinstance(cur, dict):
                cur["end_time"] = new_et
                cur["end"] = new_et
            else:
                try:
                    setattr(cur, "end_time", new_et)
                    setattr(cur, "end", new_et)
                except Exception:
                    pass
    return sorted_evs


def export_netflix_srt(events: list) -> str:
    """Generate Netflix-compliant SRT format."""
    events = _prepare_sorted_events(events)
    def format_time(seconds: float) -> str:
        if seconds is None:
            seconds = 0.0
        seconds = float(seconds)
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        if millis >= 1000:
            millis -= 1000
            secs += 1
        return f"{hrs:02d}:{mins:02d}:{secs:02d},{millis:03d}"

    entries = []
    for i, event in enumerate(events, 1):
        s_val = event.get("start_time", event.get("start", 0.0)) if isinstance(event, dict) else getattr(event, "start_time", getattr(event, "start", 0.0))
        e_val = event.get("end_time", event.get("end", 0.0)) if isinstance(event, dict) else getattr(event, "end_time", getattr(event, "end", 0.0))
        txt_val = event.get("text", "") if isinstance(event, dict) else getattr(event, "text", "")
        
        s_str = format_time(s_val)
        e_str = format_time(e_val)
        text = str(txt_val or "").replace("...", "…")
        entries.append(f"{i}\n{s_str} --> {e_str}\n{text}\n")
    
    return "\n".join(entries)


def export_netflix_vtt(events: list) -> str:
    """Generate Netflix-compliant WebVTT format."""
    events = _prepare_sorted_events(events)
    def format_time(seconds: float) -> str:
        if seconds is None:
            seconds = 0.0
        seconds = float(seconds)
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        if millis >= 1000:
            millis -= 1000
            secs += 1
        return f"{hrs:02d}:{mins:02d}:{secs:02d}.{millis:03d}"

    lines = ["WEBVTT\n"]
    for i, event in enumerate(events, 1):
        s_val = event.get("start_time", event.get("start", 0.0)) if isinstance(event, dict) else getattr(event, "start_time", getattr(event, "start", 0.0))
        e_val = event.get("end_time", event.get("end", 0.0)) if isinstance(event, dict) else getattr(event, "end_time", getattr(event, "end", 0.0))
        txt_val = event.get("text", "") if isinstance(event, dict) else getattr(event, "text", "")
        
        s_str = format_time(s_val)
        e_str = format_time(e_val)
        text = str(txt_val or "").replace("...", "…")
        lines.append(f"{i}")
        lines.append(f"{s_str} --> {e_str}")
        lines.append(f"{text}")
        lines.append("")
        
    return "\n".join(lines)


def export_netflix_ttml(events: list, language: str = "en") -> str:
    """Generate Netflix TTML/DFXP format."""
    events = _prepare_sorted_events(events)
    def format_time(seconds: float) -> str:
        if seconds is None:
            seconds = 0.0
        seconds = float(seconds)
        hrs = int(seconds // 3600)
        mins = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int(round((seconds - int(seconds)) * 1000))
        if millis >= 1000:
            millis -= 1000
            secs += 1
        return f"{hrs:02d}:{mins:02d}:{secs:02d}.{millis:03d}"

    xml = []
    xml.append('<?xml version="1.0" encoding="utf-8"?>')
    xml.append('<tt xmlns="http://www.w3.org/ns/ttml" xmlns:tts="http://www.w3.org/ns/ttml#styling" xmlns:ttm="http://www.w3.org/ns/ttml#metadata" xmlns:ttp="http://www.w3.org/ns/ttml#parameter" xml:lang="' + language + '">')
    xml.append('  <head>')
    xml.append('    <styling>')
    xml.append('      <style xml:id="default" tts:color="white" tts:fontFamily="sansSerif" tts:fontSize="100%" tts:textAlign="center" tts:origin="10% 10%" tts:extent="80% 80%"/>')
    xml.append('    </styling>')
    xml.append('    <layout>')
    xml.append('      <region xml:id="bottomCenter" style="default" tts:origin="10% 80%" tts:extent="80% 10%"/>')
    xml.append('      <region xml:id="topCenter" style="default" tts:origin="10% 10%" tts:extent="80% 10%"/>')
    xml.append('    </layout>')
    xml.append('  </head>')
    xml.append('  <body>')
    xml.append('    <div region="bottomCenter">')
    
    for event in events:
        s_val = event.get("start_time", event.get("start", 0.0)) if isinstance(event, dict) else getattr(event, "start_time", getattr(event, "start", 0.0))
        e_val = event.get("end_time", event.get("end", 0.0)) if isinstance(event, dict) else getattr(event, "end_time", getattr(event, "end", 0.0))
        txt_val = event.get("text", "") if isinstance(event, dict) else getattr(event, "text", "")
        
        s_str = format_time(s_val)
        e_str = format_time(e_val)
        text = str(txt_val or "").replace("...", "…")
        
        # Basic escaping
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        
        # Handle italics (which are now escaped)
        text = text.replace("&lt;i&gt;", '<span tts:fontStyle="italic">')
        text = text.replace("&lt;/i&gt;", '</span>')
        text = text.replace("\n", "<br/>")
            
        xml.append(f'      <p begin="{s_str}" end="{e_str}">{text}</p>')
        
    xml.append('    </div>')
    xml.append('  </body>')
    xml.append('</tt>')
    
    return "\n".join(xml)


# --- DUBBING SCRIPT ---

DUBBING_FPS = 25  # reference dubbing scripts use HH:MM:SS:FF with frames 00-24


def _dubbing_timecode(seconds: float, fps: int = DUBBING_FPS) -> str:
    """Seconds -> HH:MM:SS:FF timecode (frame-accurate, rounded to nearest frame)."""
    total_frames = max(0, int(round(float(seconds or 0.0) * fps)))
    frames = total_frames % fps
    total_secs = total_frames // fps
    return f"{total_secs // 3600:02d}:{(total_secs % 3600) // 60:02d}:{total_secs % 60:02d}:{frames:02d}"


def export_to_dubbing_script(
    result: TranscriptionResult,
    output_path: str,
    speaker_map: Optional[Dict[str, str]] = None,
) -> str:
    """Generate the Dubbing Script workbook: a dialogue sheet plus a character list sheet.

    speaker_map renames speakers (original label -> new character name); every segment
    carrying the original label gets the new name, and speakers mapped to the same name merge.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    speaker_map = {k: v.strip() for k, v in (speaker_map or {}).items() if isinstance(v, str) and v.strip()}

    def char_name(seg: Segment) -> str:
        raw = (seg.speaker or "").strip()
        return speaker_map.get(raw, raw)

    header_fill = PatternFill("solid", fgColor="00B0F0")
    header_font = Font(bold=True)
    side = Side(style="thin")
    border = Border(left=side, right=side, top=side, bottom=side)
    center = Alignment(horizontal="center", vertical="center")
    wrap = Alignment(vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = (Path(result.filename).stem or "Dubbing Script")[:31]

    tr = _has_translation(result)
    headers = ["Sr. No.", "Time In", "Time Out", "Dialogue", "Character"]
    if tr:
        label = f" ({result.translation_language})" if result.translation_language else ""
        headers.insert(4, f"Translation{label}")
    for col, title in enumerate(headers, start=2):
        c = ws.cell(row=2, column=col, value=title)
        c.fill, c.font, c.border, c.alignment = header_fill, header_font, border, center

    characters: Dict[str, Dict[str, int]] = {}
    # Consecutive lines by the same speaker collapse into one row; a new row starts only
    # when the speaker changes. Time In = first line's start, Time Out = last line's end.
    groups: List[Dict[str, Any]] = []
    for seg in sorted(result.segments, key=lambda s: s.start_time):
        text = (seg.transcript or "").strip()
        if not text:
            continue
        name = char_name(seg)
        if name:
            characters.setdefault(name, {})
            characters[name][seg.gender or "Unknown"] = characters[name].get(seg.gender or "Unknown", 0) + 1
        if groups and groups[-1]["name"] == name:
            groups[-1]["end"] = max(groups[-1]["end"], seg.end_time)
            groups[-1]["lines"].append(text)
            groups[-1]["tr"].append((seg.translation or "").strip())
        else:
            groups.append({"name": name, "start": seg.start_time, "end": seg.end_time, "lines": [text], "tr": [(seg.translation or "").strip()]})

    for i, g in enumerate(groups, start=1):
        row = i + 2
        values = [i, _dubbing_timecode(g["start"]), _dubbing_timecode(g["end"]), "\n".join(g["lines"])]
        if tr:
            values.append("\n".join(t for t in g["tr"] if t))
        values.append(g["name"] or None)
        for col, val in enumerate(values, start=2):
            c = ws.cell(row=row, column=col, value=val)
            c.border = border
            c.alignment = wrap if col in (5, 6) and (col == 5 or tr) else center

    for col, width in ({"A": 3, "B": 7, "C": 13, "D": 13, "E": 60, "F": 60, "G": 18} if tr
                       else {"A": 3, "B": 7, "C": 13, "D": 13, "E": 76, "F": 18}).items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "B3"

    ws2 = wb.create_sheet("Character List")
    for col, title in enumerate(["Sr. No.", "Character", "Gender"], start=2):
        c = ws2.cell(row=2, column=col, value=title)
        c.fill, c.font, c.border, c.alignment = header_fill, header_font, border, center
    for i, name in enumerate(sorted(characters, key=str.casefold), start=1):
        genders = characters[name]
        gender = max(genders, key=genders.get)  # most frequent gender label for that character
        for col, val in enumerate([i, name, gender], start=2):
            c = ws2.cell(row=i + 2, column=col, value=val)
            c.border, c.alignment = border, center
    for col, width in {"A": 3, "B": 7, "C": 22, "D": 10}.items():
        ws2.column_dimensions[col].width = width

    wb.save(output_path)
    return output_path
